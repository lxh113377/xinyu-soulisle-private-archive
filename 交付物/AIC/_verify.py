# -*- coding: utf-8 -*-
"""验收：原稿未改 + 派生文件合规 + 新增页背景与内容正确。"""
import hashlib
from pathlib import Path

from pptx import Presentation

ORIG = Path(r"c:\Users\37533\Desktop\workspace\项目\陪聊\交付物\提交包\心屿MindIsle_参赛方案.pptx")
HERE = Path(r"c:\Users\37533\Desktop\workspace\项目\陪聊\交付物\AIC")
DST = HERE / "心屿MindIsle_AIC参赛方案.pptx"


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()[:16]


print("=== 1) 用户原稿是否被改动 ===")
print("   当前哈希:", sha(ORIG), "(应为 91C88986B2DD2BA5)")

prs = Presentation(str(DST))
print("\n=== 2) 页数 ===")
print("   派生文件页数:", len(prs.slides), "(期望 26 = 20 用户页 + 6 新增)")

print("\n=== 3) 新增页背景是否继承（用户背景为每页图片） ===")
for i in (13, 20, 21, 22, 23):
    s = prs.slides[i]
    try:
        ft = int(s.background.fill.type)
    except Exception:
        ft = -1
    blip = s.part._element.find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}"
        "bg//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    )
    has_img_rel = any("image" in r.reltype for r in s.part.rels.values())
    print(f"   第{i+1:>2}页: bgFillType={ft} 背景blip={'有' if blip is not None else '无'} 图片关系={'有' if has_img_rel else '无'}")

print("\n=== 4) 合规扫描（AIC 公平性：不得出现校名/Logo/指导教师） ===")
blob = "\n".join(
    sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame
)
for kw in ["iCAN", "江西科技", "指导教师", "校徽", "Logo", "8 人"]:
    print(f"   {kw!r}: {blob.count(kw)}")
print("   '算法创新赛' 出现:", blob.count("算法创新赛"))

print("\n=== 5) 新增 6 页内容 ===")
for i in range(20, len(prs.slides)):
    sh = list(prs.slides[i].shapes)
    print(f"   --- 第{i+1}页 ---")
    print(f"      标签: {sh[1].text_frame.text}")
    print(f"      标题: {sh[2].text_frame.text}")
    print(f"      页码: {sh[5].text_frame.text}")
    rows = sh[4].table.rows
    print(f"      表格: {len(rows)}行 x {len(sh[4].table.columns)}列  首行: {' | '.join(c.text for c in rows[0].cells)}")
    print(f"              末行: {' | '.join(c.text[:22] for c in rows[len(rows) - 1].cells)}")
